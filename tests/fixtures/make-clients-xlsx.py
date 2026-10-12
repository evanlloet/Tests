"""
Génère tests/fixtures/clients-exemple.xlsx : un fichier clients FICTIF au même format
que l'export « CLIENTS » du laboratoire (mêmes en-têtes, dates au format Excel,
pourcentages, valeurs « - »). Aucune donnée réelle.

Usage : python3 tests/fixtures/make-clients-xlsx.py   (nécessite openpyxl)
"""
import datetime
from pathlib import Path

from openpyxl import Workbook

HEADERS = [
    " Code", "Lib Sous-Canal Client", " Nom Client", "Code Postal", "Ville", "N° Tel", "Adresse Mail ",
    "Code UGA 746 Tiers Cde", "Lib Groupement Client", "Ciblage Client Oral Care", "Ciblage Premium",
    "Ciblage Perfect Store", "Ciblage TLV", "Ciblage GC", "Nom Agent Oral Care", "KAM  régional",
    "Date de dernière Commande", "Date de prochaine visite", "CAf CFaj A-1", "CAf A", " Evol A/A-1 ",
]
D = datetime.datetime
ROWS = [
    (9000001, "PHARMACIES", "SELARL GRANDE PHIE DU PHARE", 14000, "CAEN", "+33 2 00 00 00 01", "contact@phare.example.test",
     "14CAE", "SAS GROUPEMENT EXEMPLE", "A", "-", "S", "-", "R", "Agent Exemple", "KAM Exemple", D(2026, 8, 31), D(2026, 10, 5), 100000, 140000, 0.4),
    (9000002, "GMS", "SAS EXEMPLE DISTRIBUTION", 50400, "YQUELON", "+33 2 00 00 00 02", "para@distribution.example.test",
     "50GRA", "ENSEIGNE EXEMPLE", "B", "P", "-", "-", "-", "Agent Exemple", None, D(2026, 8, 26), None, 20000, 18000, -0.1),
    (9000003, "PHARMACIES", "SELARL PHIE DES DUNES", 50110, "TOURLAVILLE", "+33 2 00 00 00 03", "dunes@example.test",
     "50CHE", None, "D", "-", "-", "-", "G", "Agent Exemple", None, D(2026, 7, 20), D(2026, 11, 19), 3000, None, None),
    (9000004, "PHARMACIES", "SELARL PHIE DE LA GARE", 95300, "HEROUVILLE", None, "gare@example.test",
     "14HER", "SAS GROUPEMENT EXEMPLE", "D", "-", "-", "-", "-", "Agent Exemple", None, D(2026, 6, 22), D(2026, 9, 14), 2800, 3136, 0.12),
    (9000005, "PHARMACIES", "MME DUPONT CLAIRE", 14220, "LE HOM", "+33 2 00 00 00 05", "dupont@example.test",
     "14VIR", "SYMBIOSE EXEMPLE", "C", "-", "-", "-", "-", "Agent Exemple", None, D(2026, 9, 1), D(2026, 10, 13), None, 66, None),
    (9000006, "PHARMACIES", "SELAS PHIE DU MARCHE", 53000, "LAVAL", "+33 2 00 00 00 06", "pas-un-mail",
     "53LAV", "SANTE EXEMPLE", "B", "-", "S", "-", "-", "Agent Exemple", None, D(2026, 8, 24), D(2026, 10, 1), 15000, 16410, 0.094),
]

wb = Workbook()
ws = wb.active
ws.title = "Feuil1"
ws.append(HEADERS)
for row in ROWS:
    ws.append(row)
for r in range(2, len(ROWS) + 2):
    ws.cell(r, 17).number_format = "mm-dd-yy"
    ws.cell(r, 18).number_format = "mm-dd-yy"
    ws.cell(r, 19).number_format = "#,##0"
    ws.cell(r, 20).number_format = "#,##0"
    ws.cell(r, 21).number_format = "0.00%"
out = Path(__file__).with_name("clients-exemple.xlsx")
wb.save(out)
print(f"{out} ({len(ROWS)} lignes)")
